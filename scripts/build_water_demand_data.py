"""Validate and encode water-demand / replenishment source data.

The source workbook holds 500k+ grid cells for one demo region. Publishing the
whole table as JSON would cost tens of megabytes, so the builder writes a
compact, deterministic binary transport plus small GeoJSON/JSON metadata:

* row-major (``gy`` then ``gx``) grid order keeps the blob highly compressible
  and lets the browser binary-search a longitude run for click hit-testing;
* longitude/latitude are grid indices into explicit axes, which removes
  floating point ambiguity from the transport;
* the six metric columns are stored as ``uint16`` quantised against their own
  observed range, which caps the absolute error well below the two decimals the
  interface displays.
"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
import gzip
import hashlib
import json
import math
import os
from pathlib import Path
import struct
import sys
import tempfile
from uuid import uuid4

import numpy as np

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook  # noqa: E402
from pyproj import CRS, Transformer
from shapely.geometry import mapping, shape
from shapely.ops import unary_union

from backend.shapefile_geojson import iter_shapefile_geojson_features  # noqa: E402

SCHEMA_VERSION = 1
UNIT = 'mm'
MAGIC = b'WDPT'
BINARY_VERSION = 1
GRID_DECIMALS = 6
QUANT_MAX = 65535
HEADER_SIZE = 120
SHAPEFILE_SIDECARS = ('.shp', '.shx', '.dbf', '.prj', '.cpg')

EXPECTED_COLUMNS = [
    'Longitude', 'Latitude', 'Class',
    'ETc', 'EWDc', 'EWRc', 'ETf', 'EWDf', 'EWRf',
]
CLASS_LEGEND = [
    {'value': 1, 'color': '#F08A85', 'label': '立即补水'},
    {'value': 2, 'color': '#F2C744', 'label': '优先补水'},
    {'value': 3, 'color': '#7BC47F', 'label': '观察复核'},
]
METRICS = [
    {'field': 'ecologicalWaterConsumption', 'label': '生态耗水', 'unit': UNIT},
    {'field': 'ecologicalWaterDemand', 'label': '生态需水', 'unit': UNIT},
    {'field': 'ecologicalWaterReplenishment', 'label': '生态补水', 'unit': UNIT},
]
VALUE_FIELDS = [
    'current.ecologicalWaterConsumption',
    'current.ecologicalWaterDemand',
    'current.ecologicalWaterReplenishment',
    'future.ecologicalWaterConsumption',
    'future.ecologicalWaterDemand',
    'future.ecologicalWaterReplenishment',
]
CLASS_VALUES = (1, 2, 3)


class BuildError(RuntimeError):
    """Raised when a source file fails a required data contract."""


def read_workbook_points(path: str | Path) -> dict[str, np.ndarray]:
    """Read and strictly validate the water-demand workbook."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.sheetnames) != 1:
            raise BuildError('workbook must contain exactly one worksheet')
        sheet = workbook[workbook.sheetnames[0]]
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        if [str(name).strip() for name in header] != EXPECTED_COLUMNS:
            raise BuildError(
                f'header must be exactly {EXPECTED_COLUMNS}, found {list(header)}'
            )

        rows: list[tuple[float, ...]] = []
        seen: set[tuple[float, float]] = set()
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if row is None or all(value is None for value in row):
                continue
            if len(row) != len(EXPECTED_COLUMNS):
                raise BuildError(f'row {row_number}: expected {len(EXPECTED_COLUMNS)} columns')
            values = _numeric_row(row, row_number)
            longitude, latitude, class_value = values[0], values[1], values[2]
            if not -180.0 <= longitude <= 180.0:
                raise BuildError(f'row {row_number}: longitude outside [-180, 180]')
            if not -90.0 <= latitude <= 90.0:
                raise BuildError(f'row {row_number}: latitude outside [-90, 90]')
            if class_value not in CLASS_VALUES:
                raise BuildError(f'row {row_number}: Class must be one of {CLASS_VALUES}')
            coordinate = (longitude, latitude)
            if coordinate in seen:
                raise BuildError(f'row {row_number}: duplicate coordinate pair')
            seen.add(coordinate)
            rows.append(values)
    finally:
        workbook.close()

    if not rows:
        raise BuildError('workbook contains no data rows')

    table = np.asarray(rows, dtype=np.float64)
    return {
        'longitude': table[:, 0],
        'latitude': table[:, 1],
        'class': table[:, 2].astype(np.uint8),
        'values': table[:, 3:9].T.copy(),
    }


def _numeric_row(row, row_number: int) -> tuple[float, ...]:
    values = []
    for column, raw in zip(EXPECTED_COLUMNS, row):
        if isinstance(raw, bool) or raw is None:
            raise BuildError(f'row {row_number}: {column} must be a finite number')
        try:
            value = float(raw)
        except (TypeError, ValueError) as exc:
            raise BuildError(f'row {row_number}: {column} must be a finite number') from exc
        if not math.isfinite(value):
            raise BuildError(f'row {row_number}: {column} must be a finite number')
        values.append(value)
    return tuple(values)


def read_region_geometry(path: str | Path) -> dict:
    """Read the demo-region Shapefile and reproject it to EPSG:4326."""
    shapefile_path = Path(path)
    prj_path = shapefile_path.with_suffix('.prj')
    if not prj_path.is_file():
        raise BuildError(f'region Shapefile is missing its .prj: {prj_path}')
    source_crs = CRS.from_wkt(prj_path.read_text(encoding='utf-8-sig'))
    transformer = Transformer.from_crs(source_crs, CRS.from_epsg(4326), always_xy=True)

    geometries = []
    for feature in iter_shapefile_geojson_features(shapefile_path):
        geometry = feature.get('geometry')
        if not geometry:
            continue
        geometries.append(_reproject_geometry(geometry, transformer))
    if not geometries:
        raise BuildError('region Shapefile contains no polygon geometry')

    merged = unary_union([shape(geometry) for geometry in geometries])
    if not merged.is_valid:
        merged = merged.buffer(0)
    if merged.geom_type not in {'Polygon', 'MultiPolygon'}:
        raise BuildError(f'region geometry resolved to unsupported {merged.geom_type}')
    return mapping(merged)


def _reproject_geometry(geometry: dict, transformer: Transformer) -> dict:
    kind = geometry.get('type')
    if kind == 'Polygon':
        return {'type': 'Polygon', 'coordinates': _reproject_rings(geometry['coordinates'], transformer)}
    if kind == 'MultiPolygon':
        return {
            'type': 'MultiPolygon',
            'coordinates': [_reproject_rings(polygon, transformer) for polygon in geometry['coordinates']],
        }
    raise BuildError(f'unsupported region geometry {kind!r}')


def _reproject_rings(rings, transformer: Transformer) -> list[list[list[float]]]:
    projected = []
    for ring in rings:
        points = []
        for coordinate in ring:
            x, y = transformer.transform(coordinate[0], coordinate[1])
            points.append([round(x, 6), round(y, 6)])
        projected.append(points)
    return projected


def load_china_outline(path: str | Path) -> dict:
    """Load a China outline geometry from a bare geometry or overview artifact."""
    payload = json.loads(Path(path).read_text(encoding='utf-8'))
    if isinstance(payload, dict) and payload.get('type') in {'Polygon', 'MultiPolygon'}:
        return payload
    outline = payload.get('chinaOutline') if isinstance(payload, dict) else None
    if isinstance(outline, dict) and outline.get('type') in {'Polygon', 'MultiPolygon'}:
        return outline
    raise BuildError(f'no China outline geometry found in {path}')


def build_grid(longitude: np.ndarray, latitude: np.ndarray) -> dict:
    """Resolve the regular sampling axes and the per-point grid indices."""
    lon_axis = np.unique(np.round(longitude, GRID_DECIMALS))
    lat_axis = np.unique(np.round(latitude, GRID_DECIMALS))
    if lon_axis.size < 2 or lat_axis.size < 2:
        raise BuildError('point grid requires at least two distinct rows and columns')

    gx = np.searchsorted(lon_axis, np.round(longitude, GRID_DECIMALS))
    gy = np.searchsorted(lat_axis, np.round(latitude, GRID_DECIMALS))
    if not np.array_equal(lon_axis[gx], np.round(longitude, GRID_DECIMALS)):
        raise BuildError('longitude values do not resolve onto the sampling axis')
    if not np.array_equal(lat_axis[gy], np.round(latitude, GRID_DECIMALS)):
        raise BuildError('latitude values do not resolve onto the sampling axis')
    if gx.max() > np.iinfo(np.uint16).max or gy.max() > np.iinfo(np.uint16).max:
        raise BuildError('sampling axis exceeds the uint16 transport limit')

    order = np.lexsort((gx, gy))
    return {
        'lon_axis': lon_axis,
        'lat_axis': lat_axis,
        'gx': gx[order].astype(np.uint16),
        'gy': gy[order].astype(np.uint16),
        'order': order,
    }


def quantize_values(values: np.ndarray) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Quantise every metric column against its own observed range."""
    if values.ndim != 2 or values.shape[0] != len(VALUE_FIELDS):
        raise BuildError(f'expected {len(VALUE_FIELDS)} metric columns')
    minimums = values.min(axis=1)
    maximums = values.max(axis=1)
    spans = maximums - minimums
    if np.any(spans <= 0):
        raise BuildError('every metric column must vary')
    scaled = (values - minimums[:, None]) / spans[:, None] * QUANT_MAX
    quantised = np.rint(scaled).astype(np.uint16)
    return quantised, minimums, maximums


def encode_points(
    point_count: int,
    class_values: np.ndarray,
    gx: np.ndarray,
    gy: np.ndarray,
    quantised: np.ndarray,
    lon_axis: np.ndarray,
    lat_axis: np.ndarray,
    minimums: np.ndarray,
    maximums: np.ndarray,
) -> bytes:
    """Pack the transport blob described by ``HEADER_SIZE``."""
    header = bytearray(HEADER_SIZE)
    struct.pack_into('<4sHHIIII', header, 0, MAGIC, BINARY_VERSION, 0, point_count,
                     lon_axis.size, lat_axis.size, quantised.shape[0])
    struct.pack_into('<6d', header, 24, *minimums.tolist())
    struct.pack_into('<6d', header, 72, *maximums.tolist())

    body = bytearray()
    body += lon_axis.astype('<f8').tobytes()
    body += lat_axis.astype('<f8').tobytes()
    body += gx.astype('<u2').tobytes()
    body += gy.astype('<u2').tobytes()
    body += np.ascontiguousarray(quantised.astype('<u2')).tobytes()
    body += class_values.astype('u1').tobytes()
    return bytes(header) + bytes(body)


def binary_offsets(point_count: int, lon_count: int, lat_count: int, column_count: int) -> dict[str, int]:
    """Return the byte offsets of each ``encode_points`` section."""
    lon_axis = HEADER_SIZE
    lat_axis = lon_axis + 8 * lon_count
    gx = lat_axis + 8 * lat_count
    gy = gx + 2 * point_count
    values = gy + 2 * point_count
    classes = values + 2 * point_count * column_count
    return {
        'lonAxis': lon_axis,
        'latAxis': lat_axis,
        'gx': gx,
        'gy': gy,
        'values': values,
        'classes': classes,
        'total': classes + point_count,
    }


def decode_points(blob: bytes) -> dict:
    """Decode a transport blob; used by the contract tests."""
    magic, version, _reserved, point_count, lon_count, lat_count, column_count = struct.unpack_from(
        '<4sHHIIII', blob, 0
    )
    if magic != MAGIC:
        raise BuildError('unexpected transport magic')
    if version != BINARY_VERSION:
        raise BuildError(f'unsupported transport version {version}')
    offsets = binary_offsets(point_count, lon_count, lat_count, column_count)
    if len(blob) != offsets['total']:
        raise BuildError('transport length does not match its header')
    minimums = np.array(struct.unpack_from('<6d', blob, 24))
    maximums = np.array(struct.unpack_from('<6d', blob, 72))
    lon_axis = np.frombuffer(blob, '<f8', lon_count, offsets['lonAxis'])
    lat_axis = np.frombuffer(blob, '<f8', lat_count, offsets['latAxis'])
    gx = np.frombuffer(blob, '<u2', point_count, offsets['gx'])
    gy = np.frombuffer(blob, '<u2', point_count, offsets['gy'])
    quantised = np.frombuffer(blob, '<u2', point_count * column_count, offsets['values']).reshape(
        column_count, point_count
    )
    classes = np.frombuffer(blob, 'u1', point_count, offsets['classes'])
    spans = maximums - minimums
    values = minimums[:, None] + quantised.astype(np.float64) * (spans[:, None] / QUANT_MAX)
    return {
        'pointCount': point_count,
        'lonAxis': lon_axis,
        'latAxis': lat_axis,
        'gx': gx,
        'gy': gy,
        'class': classes,
        'values': values,
        'minimums': minimums,
        'maximums': maximums,
    }


def region_feature(geometry: dict, name: str, region_id: str, point_count: int) -> dict:
    min_x, min_y, max_x, max_y = shape(geometry).bounds
    return {
        'type': 'Feature',
        'properties': {
            'id': region_id,
            'name': name,
            'pointCount': point_count,
            'bounds': [
                [round(min_y, 6), round(min_x, 6)],
                [round(max_y, 6), round(max_x, 6)],
            ],
        },
        'geometry': geometry,
    }


def read_region_name(path: str | Path) -> str:
    for feature in iter_shapefile_geojson_features(Path(path)):
        name = str(feature.get('properties', {}).get('name') or '').strip()
        if name:
            return name
    return '示范区域'


def encode_json(payload: object) -> bytes:
    return json.dumps(payload, ensure_ascii=False, separators=(',', ':'), sort_keys=True).encode('utf-8')


def _source_files(workbook: Path, region_shp: Path) -> list[Path]:
    files = [workbook]
    files.extend(
        region_shp.with_suffix(suffix)
        for suffix in SHAPEFILE_SIDECARS
        if region_shp.with_suffix(suffix).is_file()
    )
    return files


def _source_manifest(files: list[Path]) -> list[dict]:
    return [
        {'name': source.name, 'sha256': hashlib.sha256(source.read_bytes()).hexdigest()}
        for source in sorted(files, key=lambda item: item.name)
    ]


def _built_at(files: list[Path]) -> str:
    newest = max(path.stat().st_mtime for path in files)
    return datetime.fromtimestamp(newest, timezone.utc).isoformat(timespec='microseconds').replace('+00:00', 'Z')


def _write(path: Path, payload: bytes, *, gzip_copy: bool) -> None:
    path.write_bytes(payload)
    if gzip_copy:
        path.with_suffix(path.suffix + '.gz').write_bytes(gzip.compress(payload, compresslevel=6, mtime=0))


def build_water_demand_data(
    workbook_path: str | Path,
    region_path: str | Path,
    china_outline_path: str | Path,
    output_path: str | Path,
    *,
    force: bool = False,
) -> dict:
    """Build every deterministic offline artifact for the module."""
    workbook = Path(workbook_path)
    region_shp = Path(region_path)
    outline_path = Path(china_outline_path)
    output = Path(output_path)
    for label, source in (
        ('workbook', workbook),
        ('region Shapefile', region_shp),
        ('China outline', outline_path),
    ):
        if not source.is_file():
            raise FileNotFoundError(f'{label} not found: {source}')
    if output.exists() and not force:
        raise FileExistsError(f'output already exists; pass --force to rebuild: {output}')

    points = read_workbook_points(workbook)
    geometry = read_region_geometry(region_shp)
    outline = load_china_outline(outline_path)
    region_name = read_region_name(region_shp)
    region_id = 'WR-DEMO-1'

    grid = build_grid(points['longitude'], points['latitude'])
    ordered_values = points['values'][:, grid['order']]
    quantised, minimums, maximums = quantize_values(ordered_values)
    blob = encode_points(
        point_count=grid['gx'].size,
        class_values=points['class'][grid['order']],
        gx=grid['gx'],
        gy=grid['gy'],
        quantised=quantised,
        lon_axis=grid['lon_axis'],
        lat_axis=grid['lat_axis'],
        minimums=minimums,
        maximums=maximums,
    )

    region = region_feature(geometry, region_name, region_id, int(grid['gx'].size))
    inside = count_points_in_region(points['longitude'], points['latitude'], geometry)
    overview = {
        'schemaVersion': SCHEMA_VERSION,
        'unit': UNIT,
        'metrics': METRICS,
        'valueFields': VALUE_FIELDS,
        'legend': CLASS_LEGEND,
        'chinaOutline': outline,
        'regions': {'type': 'FeatureCollection', 'features': [region]},
    }
    manifest = {
        'schemaVersion': SCHEMA_VERSION,
        'builtAt': _built_at(_source_files(workbook, region_shp)),
        'sourceFiles': _source_manifest(_source_files(workbook, region_shp)),
        'unit': UNIT,
        'metrics': METRICS,
        'valueFields': VALUE_FIELDS,
        'legend': CLASS_LEGEND,
        'pointCount': int(grid['gx'].size),
        'pointsInsideRegion': inside,
        'pointsOutsideRegion': int(grid['gx'].size) - inside,
        'classCounts': {
            str(value): int(np.count_nonzero(points['class'] == value)) for value in CLASS_VALUES
        },
        'lonAxisCount': int(grid['lon_axis'].size),
        'latAxisCount': int(grid['lat_axis'].size),
        'binaryBytes': len(blob),
        'regions': [region['properties']],
    }

    output.parent.mkdir(parents=True, exist_ok=True)
    stage = Path(tempfile.mkdtemp(prefix=f'.{output.name}.stage-', dir=output.parent))
    try:
        _write(stage / 'region.geojson', encode_json({'type': 'FeatureCollection', 'features': [region]}), gzip_copy=False)
        _write(stage / 'overview.json', encode_json(overview), gzip_copy=True)
        _write(stage / 'points.bin', blob, gzip_copy=True)
        _write(stage / 'manifest.json', encode_json(manifest), gzip_copy=False)
        _validate_staged(stage, grid['gx'].size)
        _replace_output(stage, output, force)
    except Exception:
        if stage.exists():
            import shutil

            shutil.rmtree(stage)
        raise
    return manifest


def count_points_in_region(longitude: np.ndarray, latitude: np.ndarray, geometry: dict) -> int:
    """Count sample centres covered by the demo region."""
    from shapely.prepared import prep

    prepared = prep(shape(geometry))
    return int(sum(1 for x, y in zip(longitude, latitude) if prepared.covers(shape({'type': 'Point', 'coordinates': [x, y]}))))


def _validate_staged(stage: Path, point_count: int) -> None:
    raw = (stage / 'points.bin').read_bytes()
    if gzip.decompress((stage / 'points.bin.gz').read_bytes()) != raw:
        raise BuildError('gzip transport does not match the raw transport')
    decoded = decode_points(raw)
    if decoded['pointCount'] != point_count:
        raise BuildError('transport point count does not match the source table')
    if len(set(zip(decoded['gy'].tolist(), decoded['gx'].tolist()))) != point_count:
        raise BuildError('transport contains duplicate grid cells')
    for prefix in ('overview.json',):
        source = stage / prefix
        if gzip.decompress(source.with_suffix(source.suffix + '.gz').read_bytes()) != source.read_bytes():
            raise BuildError(f'gzip artifact does not match raw JSON: {prefix}')


def _replace_output(stage: Path, output: Path, force: bool) -> None:
    import shutil

    if output.exists() and not force:
        raise FileExistsError(f'output already exists; pass --force to rebuild: {output}')
    backup = None
    try:
        if output.exists():
            backup = output.with_name(f'.{output.name}.backup-{uuid4().hex}')
            os.replace(output, backup)
        os.replace(stage, output)
    except Exception:
        if backup is not None and backup.exists() and not output.exists():
            os.replace(backup, output)
        raise
    else:
        if backup is not None:
            shutil.rmtree(backup)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook', type=Path, required=True)
    parser.add_argument('--region-shp', type=Path, required=True)
    parser.add_argument('--china-outline', type=Path, required=True,
                        help='bare geometry or overview JSON holding chinaOutline')
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--force', action='store_true')
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    manifest = build_water_demand_data(
        args.workbook,
        args.region_shp,
        args.china_outline,
        args.output,
        force=args.force,
    )
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()


